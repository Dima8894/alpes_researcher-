#!/usr/bin/env python3
"""Package already-researched content; never fetch URLs or invent facts.

Run: python scripts/export_report.py private/report.json private/result
Dependencies: reportlab, python-docx. Font: bundled PT Sans or system DejaVu Sans.
"""
import argparse
from datetime import date
from html import escape
import json
from pathlib import Path
import re
import tempfile
from urllib.parse import urlsplit

KINDS = {"fact": "Факт", "claim": "Заявление источника", "hypothesis": "Гипотеза", "gap": "Неизвестно", "action": "Действие"}
STATUSES = {"checked", "claim", "partial", "unavailable", "not_checked", "user_material"}
STATUS_NAMES = dict(zip(["checked", "claim", "partial", "unavailable", "not_checked", "user_material"], ["Проверено", "Заявление источника", "Частично", "Недоступно", "Не проверено", "Материал пользователя"]))


def text(value, name, limit=12000, optional=False):
    if not isinstance(value, str) or len(value) > limit or (not optional and not value.strip()):
        raise ValueError(f"Некорректное поле {name}")
    if any(ord(c) < 32 and c not in '\n\t' for c in value):
        raise ValueError(f"Управляющие символы в {name}")
    return value


def iso_date(value, name):
    text(value, name, 10)
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        raise ValueError(f"{name}: нужна дата YYYY-MM-DD")
    date.fromisoformat(value)


def safe_url(value):
    text(value, "url", 3000, optional=True)
    if not value:
        return value
    u = urlsplit(value)
    if u.scheme not in ("https", "http") or not u.hostname or u.username or u.password or any(c.isspace() for c in value):
        raise ValueError("Ссылка должна быть публичным HTTP(S) URL без логина и пароля")
    return value


def validate(data):
    if not isinstance(data, dict):
        raise ValueError("Отчёт должен быть объектом JSON")
    for key in ("title", "objective", "identity"):
        text(data.get(key), key, 1500)
    iso_date(data.get("date"), "date")
    sources = data.get("sources")
    if not isinstance(sources, list) or len(sources) > 150:
        raise ValueError("sources: нужен список до 150 источников")
    ids = {}
    for s in sources:
        if not isinstance(s, dict) or not re.fullmatch(r'S[0-9]{2,3}', str(s.get("id", ""))):
            raise ValueError("Источник должен иметь ID S01…S999")
        if s['id'] in ids:
            raise ValueError("Дублирующийся ID источника")
        ids[s['id']] = s
        text(s.get('title'), 'source.title', 500)
        safe_url(s.get('url', ''))
        text(s.get('locator', ''), 'source.locator', 1000, optional=True)
        if not s.get('url') and not s.get('locator'):
            raise ValueError("У источника нужен URL или имя материала и страница")
        iso_date(s.get('checked_at'), 'source.checked_at')
        if s.get('status') not in STATUSES:
            raise ValueError("Неизвестный статус источника")
        for key in ('note', 'published_at', 'entity', 'period'):
            text(s.get(key, ''), 'source.'+key, 2000, optional=True)
    sections = data.get('sections')
    if not isinstance(sections, list) or not 1 <= len(sections) <= 12:
        raise ValueError("Нужны sections: от 1 до 12 разделов")
    groups = [data.get('summary')]
    for section in sections:
        if not isinstance(section, dict):
            raise ValueError("Раздел должен быть объектом")
        text(section.get('title'), 'section.title', 200)
        groups.append(section.get('items'))
    for items in groups:
        if not isinstance(items, list) or not 1 <= len(items) <= 80:
            raise ValueError("Нужен непустой список утверждений (до 80)")
        for item in items:
            if not isinstance(item, dict) or item.get('kind') not in KINDS:
                raise ValueError("У утверждения нужен kind: fact/claim/hypothesis/gap/action")
            text(item.get('text'), 'item.text', 6000)
            refs = item.get('sources', [])
            if not isinstance(refs, list) or any(not isinstance(r, str) or r not in ids for r in refs):
                raise ValueError("Утверждение ссылается на неизвестный источник")
            if item['kind'] in ('fact', 'claim') and not refs:
                raise ValueError("Факт или заявление должны иметь источник")
            if item['kind'] == 'fact' and any(ids[r]['status'] not in ('checked', 'user_material') for r in refs):
                raise ValueError("Факт не может опираться только на непроверенный источник")
            if item['kind'] == 'hypothesis':
                text(item.get('alternative'), 'hypothesis.alternative', 2000)
                text(item.get('question'), 'hypothesis.question', 2000)
    return data


def font_path(explicit=None):
    candidates = [Path(explicit)] if explicit else []
    candidates += [Path(__file__).resolve().parents[1] / 'assets/fonts/PTSans-Regular.ttf',
                   Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
                   Path('/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf')]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise ValueError('Не найден шрифт с кириллицей. Укажите --font с путём к подходящему TTF.')


def blocks(data):
    """One common content stream keeps PDF and DOCX in sync."""
    yield ('title', data['title'])
    yield ('meta', f"Срез: {data['date']} · {data['identity']}")
    yield ('body', 'Задача: ' + data['objective'])
    for section in [{'title': 'Главное для решения', 'items': data['summary']}, *data['sections']]:
        yield ('heading', section['title'])
        for item in section['items']:
            yield ('body', KINDS[item['kind']] + ': ' + item['text'])
            if item['kind'] == 'hypothesis':
                yield ('body', 'Альтернативное объяснение: ' + item['alternative'])
                yield ('body', 'Вопрос для проверки: ' + item['question'])
            if item.get('sources'):
                yield ('refs', item['sources'])
    yield ('heading', 'Источники и границы проверки')
    for source in data['sources']:
        yield ('body', f"{source['id']} · {source['title']} · {STATUS_NAMES[source['status']]} · проверка {source['checked_at']}")
        yield ('refs', [source['id']])
        details = [("Материал", source.get('locator')), ("Объект", source.get('entity')),
                   ("Период", source.get('period')), ("Публикация", source.get('published_at')), ("Ограничение", source.get('note'))]
        for label, value in details:
            if value:
                yield ('body', label + ': ' + value)


def write_pdf(data, path, font):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    pdfmetrics.registerFont(TTFont('Alpes', str(font)))
    styles = {
        key: ParagraphStyle(key, fontName='Alpes', fontSize=size, leading=leading,
             spaceAfter=after, spaceBefore=before, textColor=colors.HexColor(color),
             alignment=TA_LEFT, keepWithNext=key in ('title', 'heading'), splitLongWords=True)
        for key, size, leading, after, before, color in [
            ('title', 22, 27, 12, 0, '#152438'), ('heading', 15, 19, 9, 14, '#166148'),
            ('body', 11, 16, 8, 0, '#263348'), ('meta', 9, 13, 12, 0, '#526175'),
            ('refs', 9, 13, 8, 0, '#166148')]
    }
    lookup = {s['id']: s for s in data['sources']}
    story = []
    for kind, value in blocks(data):
        if kind == 'refs':
            lines = []
            for sid in value:
                s = lookup[sid]
                label = escape(sid + ' — ' + s['title'])
                lines.append(f'<link href="{escape(s["url"], quote=True)}" color="#166148">{label}</link>' if s.get('url') else label)
            markup = ' · '.join(lines)
        else:
            markup = escape(value).replace('\n', '<br/>')
        story.append(Paragraph(markup, styles[kind]))
    def page_number(canvas, doc):
        canvas.setFont('Alpes', 9)
        canvas.setFillColor(colors.HexColor('#526175'))
        canvas.drawRightString(A4[0]-56, 28, str(doc.page))
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=56, leftMargin=56,
                      topMargin=48, bottomMargin=48, title=data['title'], author='Alpes Researcher').build(
                          story, onFirstPage=page_number, onLaterPages=page_number)


def write_docx(data, path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name in ('Normal', 'Title', 'Heading 1'):
        doc.styles[name].font.name = 'Arial'
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(8)
    # Remove template-inherited paragraph borders; keep the report visually quiet.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Title'].font.color.rgb = RGBColor.from_string('152438')
    doc.styles['Heading 1'].font.size = Pt(15)
    doc.styles['Heading 1'].font.color.rgb = RGBColor.from_string('166148')
    doc.core_properties.author = 'Alpes Researcher'
    doc.core_properties.last_modified_by = 'Alpes Researcher'
    doc.core_properties.title = data['title']
    lookup = {s['id']: s for s in data['sources']}
    for kind, value in blocks(data):
        if kind == 'title': doc.add_paragraph(value, 'Title')
        elif kind == 'heading': doc.add_heading(value, 1)
        elif kind == 'refs':
            p = doc.add_paragraph()
            for i, sid in enumerate(value):
                if i: p.add_run(' · ')
                s = lookup[sid]
                label = sid + ' — ' + s['title']
                if not s.get('url'):
                    p.add_run(label)
                    continue
                link = OxmlElement('w:hyperlink')
                link.set(qn('r:id'), p.part.relate_to(s['url'], RT.HYPERLINK, is_external=True))
                run = OxmlElement('w:r')
                props = OxmlElement('w:rPr')
                color = OxmlElement('w:color'); color.set(qn('w:val'), '166148'); props.append(color)
                run.append(props)
                node = OxmlElement('w:t'); node.text = label; run.append(node)
                link.append(run); p._p.append(link)
        else: doc.add_paragraph(value)
    doc.save(path)


def write_sources(data, path):
    """Plain text source ledger. Clickable citations remain in PDF and Word."""
    parts = [data['title'], 'Источники · срез ' + data['date'], '']
    for source in data['sources']:
        parts.append(source['id'] + ' · ' + source['title'])
        if source.get('url'):
            parts.append(source['url'])
        parts.append(STATUS_NAMES[source['status']] + ' · проверка ' + source['checked_at'])
        for key, label in [('locator', 'Материал'), ('entity', 'Объект'),
                           ('published_at', 'Публикация'), ('period', 'Период'), ('note', 'Примечание')]:
            if source.get(key):
                parts.append(label + ': ' + source[key])
        parts.append('')
    path.write_text('\n'.join(parts), encoding='utf-8')


def export(data, output, font=None):
    validate(data)
    font = font_path(font)
    output = Path(output)
    # Never mix report versions or overwrite a previous deliverable implicitly.
    if output.exists():
        raise ValueError('Папка результата уже существует. Выберите новую папку версии.')
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.alpes-', dir=output.parent) as tmp:
        staging = Path(tmp)
        write_pdf(data, staging/'REPORT.pdf', font)
        write_docx(data, staging/'REPORT.docx')
        write_sources(data, staging/'SOURCES.txt')
        staging.rename(output)
    return [output / name for name in ('REPORT.pdf', 'REPORT.docx', 'SOURCES.txt')]


def main():
    parser = argparse.ArgumentParser(description='Упаковать проверенное исследование в PDF, Word и реестр источников')
    parser.add_argument('input', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--font')
    args = parser.parse_args()
    try:
        data = json.loads(args.input.read_text(encoding='utf-8'))
        for path in export(data, args.output, args.font): print(path)
    except (ValueError, OSError, ImportError) as exc:
        parser.exit(1, f'Экспорт не завершён: {exc}\n')


if __name__ == '__main__':
    main()
